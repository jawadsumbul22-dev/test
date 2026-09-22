from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\HP\Documents\Codex\2026-08-16\i\outputs\globalcommerce-enterprise-platform\docs\GlobalCommerce-Project-Report-Jawad-Ahmad.docx")
NAVY = RGBColor(16, 36, 56)
BLUE = RGBColor(46, 116, 181)
TEAL = RGBColor(10, 132, 125)
GOLD = RGBColor(174, 128, 31)
MUTED = RGBColor(96, 112, 124)
LIGHT = "F2F4F7"
PALE_TEAL = "E5F3F1"
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_inches):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total = int(sum(widths_inches) * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_inches:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            dxa = int(widths_inches[idx] * 1440)
            cell.width = Inches(widths_inches[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(dxa))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, end])


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    repeat_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = NAVY
    for row_data in rows:
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for idx, value in enumerate(row_data):
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(value))
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.font.color.rgb = NAVY
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_TEAL)
    p_pr.append(shd)
    r1 = p.add_run(f"{label}: ")
    r1.bold = True
    r1.font.color.rgb = TEAL
    r2 = p.add_run(text)
    r2.font.color.rgb = NAVY


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = NAVY
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.333

for name, size, color, before, after in [
    ("Title", 30, NAVY, 0, 8),
    ("Subtitle", 14, MUTED, 0, 16),
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("GLOBALCOMMERCE  /  ENTERPRISE CAPSTONE PROJECT")
hr.bold = True
hr.font.name = "Calibri"
hr.font.size = Pt(8.5)
hr.font.color.rgb = MUTED
page_number(section.footer.paragraphs[0])

# Editorial cover.
for _ in range(4):
    doc.add_paragraph()
kicker = doc.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
kicker.paragraph_format.space_after = Pt(18)
kr = kicker.add_run("EDUQUAL LEVEL 6  |  DIPLOMA IN ARTIFICIAL INTELLIGENCE OPERATIONS")
kr.bold = True
kr.font.name = "Calibri"
kr.font.size = Pt(10)
kr.font.color.rgb = GOLD
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("GlobalCommerce Enterprise Platform")
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Marketplace management, intelligent inventory, international fulfilment, logistics visibility, AI-assisted demand forecasting, and executive commerce analytics")
author = doc.add_paragraph()
author.alignment = WD_ALIGN_PARAGRAPH.CENTER
author.paragraph_format.space_before = Pt(34)
author.paragraph_format.space_after = Pt(4)
ar = author.add_run("Jawad Ahmad")
ar.bold = True
ar.font.size = Pt(13)
ar.font.color.rgb = NAVY
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
mr = meta.add_run("Enterprise Capstone Project 21  |  Local proof of concept  |  August 2026")
mr.font.size = Pt(10)
mr.font.color.rgb = MUTED
doc.add_page_break()

doc.add_heading("Document control", level=1)
add_table(doc, ["Field", "Value"], [
    ("Student", "Jawad Ahmad"),
    ("Qualification", "EduQual Level 6 Diploma in Artificial Intelligence Operations"),
    ("Assessment client", "GlobalCommerce Marketplace International (fictional)"),
    ("Environment", "Windows 10, 16 GB RAM, Docker Compose, simulated data"),
    ("Status", "Local build; not published to GitHub or hosted externally"),
], [1.6, 4.9])

doc.add_heading("Executive summary", level=1)
doc.add_paragraph("GlobalCommerce Marketplace International is a fictional fast-growing cross-border marketplace whose merchant, inventory, payment, order, logistics, and analytics operations are fragmented across disconnected systems. This project designs and implements a laptop-efficient proof of concept that unifies those operations and adds AI-assisted inventory planning and executive visibility.")
doc.add_paragraph("The implemented solution uses React, FastAPI, PostgreSQL, Redis, scikit-learn, Docker Compose, Prometheus, and Grafana. It demonstrates a complete international order from merchant and product master data through warehouse reservation, simulated landed cost and payment, shipment tracking, demand forecasting, replenishment advice, and executive KPIs.")
add_callout(doc, "Central conclusion", "A modular, observable commerce platform can provide a defensible end-to-end operating picture on a 16 GB assessment laptop while preserving a credible path to enterprise scale.")

doc.add_heading("1. Business context and problem", level=1)
doc.add_paragraph("Disconnected commerce systems produce inconsistent master data, stock uncertainty, slow fulfilment, overstocking, shortages, weak merchant governance, limited tracking, and delayed executive decisions. Cross-border trade adds currency, duty, tax, warehouse, logistics, privacy, and regulatory complexity.")
for text in [
    "Poor visibility of available versus reserved inventory across regional warehouses.",
    "Delayed fulfilment and inconsistent customer tracking.",
    "No shared merchant verification and operational audit trail.",
    "Limited use of historical demand when planning replenishment.",
    "Executive reports that are delayed or disconnected from transactions.",
]:
    bullet(doc, text)

doc.add_heading("2. Objectives and scope", level=1)
objectives = [
    ("Marketplace", "Centralize merchant verification, customers, products, and marketplace statistics."),
    ("Inventory", "Track multi-warehouse stock, reservations, movements, health, and utilization."),
    ("International orders", "Demonstrate idempotent ordering, warehouse selection, landed cost, and simulated payment."),
    ("Logistics", "Create shipments and expose normal and delayed tracking states."),
    ("AI planning", "Evaluate lightweight forecasts and generate explainable replenishment advice."),
    ("Executive intelligence", "Update KPIs directly from persisted operational state."),
    ("Engineering", "Deploy reproducibly with security, tests, documentation, and monitoring."),
]
add_table(doc, ["Objective", "Success evidence"], objectives, [1.55, 4.95])
doc.add_paragraph("Only simulated customers, products, payments, exchange rates, duties, tax, logistics, inventory, and sales histories are used. The proof of concept demonstrates the architecture for five-million-transaction scale; it does not claim to execute that production load on a laptop.")

doc.add_heading("3. Solution architecture", level=1)
doc.add_paragraph("The solution uses a modular monolith: one FastAPI deployment with explicit marketplace, inventory, order/logistics, forecasting, and analytics responsibilities. PostgreSQL is authoritative, Redis supports rate limiting, and Prometheus/Grafana provide operational evidence.")
add_table(doc, ["Component", "Responsibility", "Reason"], [
    ("React console", "Authenticated operations and executive views", "Clear live demonstration and responsive workflow"),
    ("FastAPI", "REST contracts, roles, transactions, forecasting and KPIs", "Python alignment, OpenAPI and low laptop overhead"),
    ("PostgreSQL", "Commerce, inventory, audit and model-result state", "Relational integrity and transactional locking"),
    ("Redis", "Rate-limit counters and future short-lived coordination", "Fast ephemeral control without becoming source of truth"),
    ("Prometheus/Grafana", "Request rate, latency, status and operations evidence", "Employer-style observability"),
], [1.25, 3.1, 2.15])
add_callout(doc, "Architecture trade-off", "Microservices would improve independent scaling but add distributed transactions, service discovery, tracing, deployment, and failure modes. The modular monolith is intentionally optimized for the assessment constraint and can be separated when evidence justifies it.")

doc.add_heading("4. Marketplace and merchant management", level=1)
doc.add_paragraph("The marketplace module manages authenticated roles, merchant organizations, verification status, customer profiles, products, and categories. Administrative review creates audit evidence rather than silently changing seller state.")
for text in [
    "Merchant states: pending, verified, and suspended.",
    "Customer country and preferred currency support international checkout.",
    "Products carry SKU, merchant, category, USD base price, weight, and HS-like classification.",
    "Dashboard statistics are calculated from durable records.",
]:
    bullet(doc, text)

doc.add_heading("5. Inventory and warehouse management", level=1)
doc.add_paragraph("Each product/warehouse balance stores on-hand, reserved, and reorder-point quantities. Available stock is computed as on hand minus reserved. Receipts, reservations, releases, shipments, and adjustments are stored in a movement ledger.")
add_callout(doc, "Concurrency control", "Order creation locks candidate balance rows and reserves stock inside the same database transaction. If another request wins the final units, the losing request receives a conflict and creates no partial order, payment, or shipment.")

doc.add_heading("6. International order and logistics implementation", level=1)
for text in [
    "Check the unique idempotency key and return an existing order for a retry.",
    "Validate the customer, products, quantity, currency, and availability.",
    "Select the warehouse with the highest available quantity and reserve stock.",
    "Convert the USD base amount to the selected simulated currency.",
    "Calculate simulated shipping, five-percent duty, and eight-percent tax.",
    "Authorize a simulated payment and create a tracking number.",
    "Write the initial logistics event and audit record before commit.",
]:
    numbered(doc, text)
doc.add_paragraph("Shipments progress through created, picked, in transit, customs clearance, out for delivery, delivered, or delayed states. Delivery captures the simulated payment and deducts the reserved physical stock.")

doc.add_heading("7. AI-assisted demand forecasting", level=1)
doc.add_paragraph("Each product/warehouse series contains 120 days of deterministic simulated sales with trend, weekly seasonality, and controlled noise. The model uses lag 1, lag 7, lag 14, shifted rolling seven-day mean, weekday, and trend features.")
add_table(doc, ["Design element", "Implementation", "Risk control"], [
    ("Split", "Chronological 80/20 holdout", "Prevents future-to-past leakage"),
    ("Model", "Random Forest, 80 trees, depth 6, one CPU thread", "Lightweight and reproducible"),
    ("Metrics", "MAE and RMSE", "Measures typical and larger errors"),
    ("Interval", "Prediction +/- 1.96 residual standard deviations", "Discloses uncertainty; not called guaranteed confidence"),
    ("Decision", "Forecast + safety stock - available inventory", "Transparent rationale and human approval"),
], [1.2, 3.15, 2.15])
doc.add_paragraph("The model is decision support, not an autonomous purchasing agent. Production adoption requires baseline comparison, rolling-origin validation, drift monitoring, ownership, versioning, approval criteria, and rollback.")

doc.add_heading("8. Executive commerce intelligence", level=1)
doc.add_paragraph("The executive endpoint calculates verified merchants, customers, products, warehouses, order count, in-motion orders, inventory units, low-stock positions, GMV, delivery performance, recent orders, and replenishment risk from the current database state. The dashboard therefore changes after the live transaction instead of displaying fixed presentation data.")

doc.add_heading("9. Security, governance, and standards", level=1)
controls = [
    ("Identity and access", "JWT, Argon2, active-user check, server-side role enforcement", "ISO 27001; NIST CSF; OWASP"),
    ("Transaction integrity", "Row locks, idempotency, unique identifiers, rollback", "ISO 9001; OWASP API"),
    ("Data quality", "Schemas, master-data uniqueness, traceable corrections", "ISO 8000; GS1 consideration"),
    ("Supply-chain security", "Tracking events, exceptions, audit and chain visibility", "ISO 28000"),
    ("AI governance", "Purpose, metrics, interval, limitations, human validation", "ISO 42001; NIST AI RMF"),
    ("Operations", "Metrics, logs, request IDs, scans and recovery scenarios", "ITIL 4; COBIT 2019; CIS Controls"),
]
add_table(doc, ["Control area", "Evidence", "Framework connection"], controls, [1.25, 3.35, 1.9])

doc.add_heading("10. Testing and operational evidence", level=1)
doc.add_paragraph("Automated tests cover authentication, protected endpoints, international order creation, inventory reservation, payment authorization/capture, shipment delivery, dashboard updates, idempotency, and forecasting. The GitHub-ready workflow also defines frontend type checks, builds, Semgrep, and Trivy, but no repository publication occurs in this local build.")
for text in [
    "Health and database-readiness endpoints.",
    "Prometheus request count and latency histograms.",
    "Provisioned Grafana panels for request rate, p95 latency, and statuses.",
    "Numbered Postman collection for the live workflow.",
    "Deterministic reset script for reliable demonstration recovery.",
]:
    bullet(doc, text)

doc.add_heading("11. Results and business value", level=1)
doc.add_paragraph("The project produces a coherent proof of concept in which one business transaction updates inventory, payment, logistics, audit, forecasting inputs, and executive views. This directly demonstrates improved operational visibility, traceability, customer communication, inventory planning, and management decision support.")
add_callout(doc, "Level 6 evidence", "The strongest assessment evidence is not the number of technologies. It is the ability to justify architecture, protect transaction integrity, interpret AI uncertainty, connect controls to business risk, and explain how the laptop proof of concept evolves toward enterprise scale.")

doc.add_heading("12. Limitations", level=1)
for text in [
    "Simulated payment, exchange, customs, duty, tax, carrier, and customer data.",
    "No split shipment, back-order, returns, refund, or procurement workflow.",
    "No managed identity provider, MFA, or final customer object-ownership policy.",
    "No event broker, multi-region database, or production backup service.",
    "Small synthetic forecasting dataset and approximate prediction intervals.",
    "Local HTTP and monitoring credentials are assessment conveniences only.",
]:
    bullet(doc, text)

doc.add_heading("13. Recommendations and future enhancements", level=1)
for text in [
    "Add external identity, MFA, per-object authorization, and privacy lifecycle controls.",
    "Integrate real payment intents, carrier webhooks, currency, and customs adapters.",
    "Introduce a transactional outbox, message broker, and asynchronous consumers.",
    "Add split fulfilment, returns, refunds, procurement, and supplier lead times.",
    "Use rolling-origin validation, baseline comparison, drift monitoring, and a model registry.",
    "Implement regional resilience, encrypted backups, tracing, alerting, and performance certification.",
    "Adopt governed GS1 identifiers and master-data workflows.",
]:
    numbered(doc, text)

doc.add_heading("14. Mandatory architecture and engineering deliverables", level=1)
diagram_rows = [
    (1, "Enterprise digital commerce architecture"), (2, "High-level solution architecture"),
    (3, "Detailed system architecture"), (4, "Marketplace management architecture"),
    (5, "Inventory management architecture"), (6, "Warehouse management architecture"),
    (7, "Logistics management architecture"), (8, "AI demand forecasting architecture"),
    (9, "Enterprise network architecture"), (10, "Network flow"),
    (11, "Data flow Level 0"), (12, "Data flow Level 1"),
    (13, "Marketplace workflow"), (14, "Order fulfilment workflow"),
    (15, "Demand forecasting workflow"), (16, "Logistics workflow"),
    (17, "Database ERD"), (18, "End-to-end sequence diagram"),
]
add_table(doc, ["No.", "Editable diagram source"], diagram_rows, [0.65, 5.85])
doc.add_paragraph("All diagram sources are stored in the repository's `diagrams` directory and are designed for presentation and live redesign discussion.")

doc.add_heading("15. Conclusion", level=1)
doc.add_paragraph("GlobalCommerce demonstrates strategic architecture, risk-based trade-offs, integrated implementation, AI governance, enterprise operations, professional evidence, and a defensible path from a laptop proof of concept to a global marketplace. The system is intentionally local and simulated, but its business flow and engineering controls are real, connected, testable, and explainable.")

doc.add_heading("References and assessment basis", level=1)
for text in [
    "Al Nafi International College. Examination Readiness & Presentation Compliance Instructions. 9 January 2026.",
    "Al Nafi International College. Exam Topic for Oral Presentation - Enterprise Capstone Project 21. 15 August 2026.",
    "NIST. Artificial Intelligence Risk Management Framework.",
    "NIST. Cybersecurity Framework 2.0.",
    "OWASP Foundation. OWASP Top 10 and OWASP API Security Top 10.",
    "ISO/IEC 27001, 27017, 27018, and 42001; ISO 8000, 9001, and 28000; GS1 standards; CIS Controls; COBIT 2019; ITIL 4.",
]:
    bullet(doc, text)

doc.core_properties.title = "GlobalCommerce Enterprise Platform - Project Report"
doc.core_properties.subject = "EduQual Level 6 Enterprise Capstone Project"
doc.core_properties.author = "Jawad Ahmad"
doc.core_properties.keywords = "EduQual, AIOps, cross-border commerce, FastAPI, forecasting, Docker"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
