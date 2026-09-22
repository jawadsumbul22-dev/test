from pathlib import Path
from docx import Document

path = Path(r"C:\Users\HP\Documents\Codex\2026-08-16\i\outputs\globalcommerce-enterprise-platform\docs\GlobalCommerce-Project-Report-Jawad-Ahmad.docx")
doc = Document(path)
print({
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "sections": len(doc.sections),
    "bytes": path.stat().st_size,
    "headings": sum(1 for paragraph in doc.paragraphs if paragraph.style and paragraph.style.name.startswith("Heading")),
    "empty_table_cells": sum(1 for table in doc.tables for row in table.rows for cell in row.cells if not cell.text.strip()),
})
for index, table in enumerate(doc.tables, 1):
    widths = [round(cell.width.inches, 3) if cell.width else None for cell in table.rows[0].cells]
    print({"table": index, "rows": len(table.rows), "columns": len(table.columns), "widths": widths})
