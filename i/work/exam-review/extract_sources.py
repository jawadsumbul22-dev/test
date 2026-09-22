from pathlib import Path
import json
import pdfplumber
from pptx import Presentation

WORK = Path(r"C:\Users\HP\Documents\Codex\2026-08-16\i\work\exam-review")
PDFS = {
    "exam-guide": Path(r"C:\Users\HP\Downloads\Oral Presentation Exam Guide.pdf"),
    "topic": Path(r"C:\Users\HP\Downloads\Jawad Ahmad _Exam Topic for Oral Presentation_DAIOL6.pdf"),
}
PPTX = Path(r"C:\Users\HP\Downloads\Sample Presentation EduQual Diplomas (1).pptx")

summary = {"pdfs": {}, "presentation": {}}

for key, path in PDFS.items():
    pages = []
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            pages.append({
                "page": number,
                "width": page.width,
                "height": page.height,
                "text": page.extract_text(x_tolerance=2, y_tolerance=2) or "",
                "tables": page.extract_tables(),
            })
    summary["pdfs"][key] = {"path": str(path), "page_count": len(pages), "pages": pages}
    text = "\n\n".join(f"===== PAGE {p['page']} =====\n{p['text']}" for p in pages)
    (WORK / f"{key}-extracted.txt").write_text(text, encoding="utf-8")

prs = Presentation(PPTX)
slides = []
for index, slide in enumerate(prs.slides, 1):
    items = []
    for shape in slide.shapes:
        item = {
            "name": shape.name,
            "shape_type": str(shape.shape_type),
            "left": int(shape.left),
            "top": int(shape.top),
            "width": int(shape.width),
            "height": int(shape.height),
        }
        if hasattr(shape, "text") and shape.text.strip():
            item["text"] = shape.text
        if getattr(shape, "has_table", False):
            item["table"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
        items.append(item)
    notes = ""
    try:
        notes = slide.notes_slide.notes_text_frame.text
    except Exception:
        pass
    slides.append({"slide": index, "layout": slide.slide_layout.name, "items": items, "notes": notes})

summary["presentation"] = {
    "path": str(PPTX),
    "slide_count": len(slides),
    "slide_width": int(prs.slide_width),
    "slide_height": int(prs.slide_height),
    "slides": slides,
}
(WORK / "sources.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")

lines = []
for slide in slides:
    lines.append(f"===== SLIDE {slide['slide']} | {slide['layout']} =====")
    for item in slide["items"]:
        if "text" in item:
            lines.append(f"[{item['name']}] {item['text']}")
        if "table" in item:
            lines.append(f"[{item['name']} TABLE] {item['table']}")
    if slide["notes"].strip():
        lines.append(f"[NOTES] {slide['notes']}")
    lines.append("")
(WORK / "presentation-extracted.txt").write_text("\n".join(lines), encoding="utf-8")

print(json.dumps({
    "pdf_pages": {k: v["page_count"] for k, v in summary["pdfs"].items()},
    "ppt_slides": len(slides),
}, indent=2))
