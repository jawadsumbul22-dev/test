from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader
import json
ROOT = Path(r"C:\Users\HP\Documents\Codex\2026-08-16\i")
PROJECT = ROOT/"outputs/globalcommerce-enterprise-platform"
OUT = PROJECT/"docs"
source = (OUT/"project-report.md").read_text(encoding="utf-8")
chapters = source.split("\n## ")[1:]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.8)
sec.left_margin = sec.right_margin = Inches(.85)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.16
normal.paragraph_format.space_after = Pt(10)
normal.font.color.rgb = RGBColor.from_string("203044")
heading = doc.styles["Heading 1"]
heading.font.name = "Calibri"; heading.font.size = Pt(21)
heading.font.color.rgb = RGBColor.from_string("164D67")
heading.paragraph_format.space_after = Pt(18)
heading.paragraph_format.keep_with_next = True
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run("GlobalCommerce | Jawad Ahmad | ")
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
for run in footer.runs: run.font.size = Pt(9)
doc.add_paragraph("EDUQUAL LEVEL 6", "Subtitle")
doc.add_paragraph("GlobalCommerce\nEnterprise Platform", "Title")
doc.add_paragraph("Project report", "Subtitle")
doc.add_paragraph("Jawad Ahmad\nDiploma in Artificial Intelligence Operations\n6 September 2026")
doc.add_paragraph("A local cross-border commerce proof of concept", "Heading 2")
doc.add_paragraph("Marketplace management, inventory optimization, AI-assisted demand forecasting, international orders, logistics visibility and executive analytics.")
doc.add_paragraph("Assessment edition. External commerce providers and seed sales data are simulated. Source code and submission remain local.")
for chapter in chapters:
    title, body = chapter.split("\n", 1)
    doc.add_page_break()
    doc.add_heading(title, 1)
    for paragraph in body.strip().split("\n\n"):
        doc.add_paragraph(paragraph.strip())
doc.core_properties.title = "GlobalCommerce - Final Project Report"
doc.core_properties.author = "Jawad Ahmad"
doc.core_properties.subject = "EduQual Level 6 local assessment project"
docx = OUT/"GlobalCommerce-Final-Report-Jawad-Ahmad.docx"
doc.save(docx)

pdfmetrics.registerFont(TTFont("Calibri", "C:/Windows/Fonts/calibri.ttf"))
pdfmetrics.registerFont(TTFont("CalibriBold", "C:/Windows/Fonts/calibrib.ttf"))
body_style = ParagraphStyle("Body", fontName="Calibri", fontSize=11, leading=14.7, textColor=HexColor("#203044"), spaceAfter=11)
title_style = ParagraphStyle("Chapter", fontName="CalibriBold", fontSize=21, leading=26, textColor=HexColor("#164D67"), spaceAfter=19)
cover_style = ParagraphStyle("Cover", fontName="CalibriBold", fontSize=34, leading=39, textColor=HexColor("#164D67"), spaceAfter=24)
subtitle = ParagraphStyle("Sub", fontName="Calibri", fontSize=16, leading=21, textColor=HexColor("#465568"), spaceAfter=22)
story = [Spacer(1,55), Paragraph("EDUQUAL LEVEL 6", subtitle), Paragraph("GlobalCommerce<br/>Enterprise Platform", cover_style),
         Paragraph("Project report", subtitle), Spacer(1,22), Paragraph("Jawad Ahmad<br/>Diploma in Artificial Intelligence Operations<br/>6 September 2026", subtitle),
         Spacer(1,24), Paragraph("A local cross-border commerce proof of concept", title_style),
         Paragraph("Marketplace management, inventory optimization, AI-assisted demand forecasting, international orders, logistics visibility and executive analytics.", body_style),
         Spacer(1,15), Paragraph("Assessment edition. External commerce providers and seed sales data are simulated. Source code and submission remain local.", body_style)]
for chapter in chapters:
    title, body = chapter.split("\n", 1)
    story += [PageBreak(), Paragraph(escape(title),title_style)]
    story += [Paragraph(escape(p.strip()),body_style) for p in body.strip().split("\n\n")]
def footer_page(canvas, document):
    canvas.saveState()
    canvas.setFont("Calibri",9); canvas.setFillColor(HexColor("#617185"))
    canvas.drawString(51,28,"GlobalCommerce | Jawad Ahmad | 6 September 2026")
    canvas.drawRightString(A4[0]-51,28,str(document.page))
    canvas.restoreState()
pdf = OUT/"GlobalCommerce-Final-Report-Jawad-Ahmad.pdf"
SimpleDocTemplate(str(pdf),pagesize=A4,leftMargin=61,rightMargin=61,topMargin=54,bottomMargin=51,
                  title="GlobalCommerce - Final Project Report",author="Jawad Ahmad").build(story,onFirstPage=footer_page,onLaterPages=footer_page)
pages = PdfReader(pdf).pages
checks = dict(pdf_pages=len(pages),chapters=len(chapters),docx_paragraphs=len(doc.paragraphs),
              pdf_text_characters=sum(len(p.extract_text()) for p in pages),
              word_layout="Not rendered: bundled LibreOffice absent on this Windows runtime. PDF independently rendered and reviewed.")
(ROOT/"work/report-final").mkdir(exist_ok=True)
(ROOT/"work/report-final/structure.json").write_text(json.dumps(checks,indent=2))
print(json.dumps(checks))
